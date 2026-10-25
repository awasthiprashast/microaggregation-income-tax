"""
Export docs/REPORT.md to a formatted Word document (and PDF) for submission.

    python scripts/export_report.py --name "Your Name" --roll "123" --course "Data Privacy" \
        --instructor "Dr. X" --institution "Your University"

Produces docs/Project_Report.docx and, on Windows with Microsoft Word installed,
docs/Project_Report.pdf (Word also refreshes the table of contents).
Requires: pip install python-docx
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ACCENT = RGBColor(0x1C, 0x4E, 0x8A)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*?\*)")


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, TOC, ...) as a complex field."""
    def el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e

    run = paragraph.add_run()
    run._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = el("w:instrText", **{"xml:space": "preserve"})
    instr.text = instruction
    run._r.append(instr)
    run._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    paragraph.add_run(placeholder)
    end = paragraph.add_run()
    end._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_inline(paragraph, text: str, size: float | None = None) -> None:
    """Add text with **bold**, *italic*, `code` and [link](url) (link rendered as text)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            run = paragraph.add_run(part[1:part.index("](")])
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def full_width(table) -> None:
    """Make a table span the full text width (Word's 'preferred width' = 100 %)."""
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")


# --------------------------------------------------------------------------- #
# Document set-up
# --------------------------------------------------------------------------- #
def setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)


def add_footer(doc, text: str) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{text}   ·   Page ")
    run.font.size = Pt(9)
    add_field(p, "PAGE", "1")
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def cover_page(doc, args, title: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(args.course.upper() if args.course else "DATA PRIVACY")
    r.font.size, r.bold, r.font.color.rgb = Pt(12), True, RGBColor(0x59, 0x59, 0x59)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size, r.bold, r.font.color.rgb = Pt(24), True, ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(args.doctype)
    r.font.size = Pt(15)
    for _ in range(4):
        doc.add_paragraph()

    rows = [("Submitted by", args.name), ("Roll / Registration No.", args.roll),
            ("Course", args.course), ("Submitted to", args.instructor),
            ("Institution", args.institution), ("Date", args.date)]
    rows = [(k, v) for k, v in rows if v]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        a, b = table.rows[i].cells
        a.width, b.width = Cm(5.5), Cm(8.5)
        ra = a.paragraphs[0].add_run(k)
        ra.bold, ra.font.size = True, Pt(12)
        rb = b.paragraphs[0].add_run(v)
        rb.font.size = Pt(12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if args.no_toc:
        return
    p = doc.add_paragraph()
    r = p.add_run("Table of Contents")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), ACCENT
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              "Right-click here and choose 'Update Field' to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- #
# Markdown -> docx
# --------------------------------------------------------------------------- #
def add_table(doc, lines: list[str], counter: dict) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]   # drop |---| row
    header, body = rows[0], rows[1:]
    n_cols = len(header)
    size = 7 if n_cols >= 12 else 7.5 if n_cols >= 9 else 8.5 if n_cols >= 7 else 9.5

    table = doc.add_table(rows=1 + len(body), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    full_width(table)
    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        shade(cell, "DCE6F1")
        p = cell.paragraphs[0]
        add_inline(p, text, size)
        for r in p.runs:
            r.bold = True
    for i, row in enumerate(body, start=1):
        for j in range(n_cols):
            cell = table.rows[i].cells[j]
            text = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            add_inline(p, text, size)
            if re.fullmatch(r"[-–+]?[\d.,]+%?", text):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i % 2 == 0:
                shade(cell, "F5F7FA")
    for row in table.rows:                                   # repeat header row, tight cells
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    counter["table"] += 1
    doc.add_paragraph()


def add_figure(doc, alt: str, src: str, counter: dict) -> None:
    path = (DOCS / src).resolve()
    if not path.exists():
        doc.add_paragraph(f"[missing figure: {src}]")
        return
    sec = doc.sections[-1]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=int(width * 0.95))
    counter["figure"] += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Figure {counter['figure']}. ")
    r.bold, r.font.size = True, Pt(9.5)
    r = cap.add_run(alt)
    r.italic, r.font.size = True, Pt(9.5)


def add_code(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade(cell, "F2F2F2")
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        r.font.name, r.font.size = "Consolas", Pt(8)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph()


def convert(md: str, doc) -> None:
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    lines = md.splitlines()
    counter = {"figure": 0, "table": 0}
    i, para = 0, []

    def flush():
        if para:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            flush()
            i += 1
            continue
        if s.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue
        if s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, counter)
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            flush()
            add_figure(doc, m.group(1), m.group(2), counter)
            i += 1
            continue
        m = re.match(r"(#{1,4})\s+(.*)", s)
        if m:
            flush()
            level = len(m.group(1))
            if level > 1:                      # level 1 = document title (on the cover page)
                doc.add_heading(m.group(2), level=level - 1)
            i += 1
            continue
        if s.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, " ".join(quote), 10)
            for r in p.runs:
                r.italic = True
            continue
        m = re.match(r"([*-]|\d+\.)\s+(.*)", s)
        if m and not line.startswith("  "):
            flush()
            item = [m.group(2)]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                item.append(lines[i].strip())        # indented continuation line
                i += 1
            if m.group(1)[0].isdigit():
                # explicit numbers: Word's "List Number" style would continue numbering
                # across separate lists in the document
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.add_run(f"{m.group(1)}\t")
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
            else:
                p = doc.add_paragraph(style="List Bullet")
            add_inline(p, " ".join(item))
            continue
        para.append(line)
        i += 1
    flush()


# --------------------------------------------------------------------------- #
# PDF via Microsoft Word (Windows)
# --------------------------------------------------------------------------- #
def word_to_pdf(docx_path: Path, pdf_path: Path) -> bool:
    ps = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
  $doc = $word.Documents.Open('{docx_path}')
  foreach ($toc in $doc.TablesOfContents) {{ $toc.Update() }}
  $doc.Fields.Update() | Out-Null
  $doc.Save()
  $doc.ExportAsFixedFormat('{pdf_path}', 17)
  $doc.Close()
}} finally {{ $word.Quit() }}
"""
    try:
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                       check=True, timeout=240, capture_output=True, text=True)
        return pdf_path.exists()
    except Exception as exc:  # Word missing, not Windows, timeout...
        print(f"PDF export skipped ({exc}). Open the .docx in Word and use File > Save As > PDF.")
        return False


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--name", default="[Your Name]")
    ap.add_argument("--roll", default="[Roll / Registration No.]")
    ap.add_argument("--course", default="Data Privacy")
    ap.add_argument("--instructor", default="[Instructor Name]")
    ap.add_argument("--institution", default="[Institution]")
    ap.add_argument("--date", default=date.today().strftime("%d %B %Y"))
    ap.add_argument("--out", default="Project_Report")
    ap.add_argument("--src", default="REPORT.md", help="Markdown file inside docs/")
    ap.add_argument("--doctype", default="Project Report", help="subtitle on the cover page")
    ap.add_argument("--no-toc", action="store_true")
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args(argv)

    md = (DOCS / args.src).read_text(encoding="utf-8")
    title = re.search(r"^#\s+(.*)$", md, flags=re.M).group(1)
    md = re.sub(r"^\*\*Project report\*\*.*$", "", md, flags=re.M)     # replaced by the cover page

    doc = Document()
    setup_styles(doc)
    add_footer(doc, "Micro-aggregation in Income and Tax Data Publication")  # footer text
    cover_page(doc, args, title)
    convert(md, doc)

    docx_path = DOCS / f"{args.out}.docx"
    doc.save(docx_path)
    print(f"Saved {docx_path}")
    if not args.no_pdf and sys.platform == "win32":
        pdf_path = DOCS / f"{args.out}.pdf"
        if word_to_pdf(docx_path, pdf_path):
            print(f"Saved {pdf_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
